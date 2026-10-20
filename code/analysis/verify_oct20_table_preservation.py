#!/usr/bin/env python3
"""Check exact Word table-cell text preservation across Oct 17 -> Oct 20.
No inference, training, or numerical truth verification occurs.
Usage: python code/analysis/verify_oct20_table_preservation.py OLD_MAIN.docx NEW_MAIN.docx OLD_SUPP.docx NEW_SUPP.docx
"""
import csv,hashlib,json,sys
from pathlib import Path
from docx import Document
if len(sys.argv)!=5:raise SystemExit(__doc__)
labels=("MAIN","SUPPLEMENT");rows=[];docs=[];errors=[]
for idx,label in enumerate(labels):
 p,q=map(Path,(sys.argv[1+idx*2],sys.argv[2+idx*2]))
 a,b=Document(p),Document(q)
 digest=lambda x:hashlib.sha256(x.read_bytes()).hexdigest()
 docs.append(dict(document=label,old_sha256=digest(p),new_sha256=digest(q),old_tables=len(a.tables),new_tables=len(b.tables)))
 if len(a.tables)!=len(b.tables):errors.append(label+": table count differs")
 for j,(x,y) in enumerate(zip(a.tables,b.tables),1):
  left=[[c.text for c in r.cells] for r in x.rows]
  right=[[c.text for c in r.cells] for r in y.rows]
  good=left==right
  rows.append(dict(document=label,table=j,cell_text_preserved=good))
  if not good:errors.append(f"{label} table {j}")
result=dict(status="PASS_TABLE_TEXT_ONLY" if not errors else "FAIL",tables_checked=len(rows),mismatches=errors,documents=docs,table_decisions=rows,
 limitations="Does not recompute numerical paper tables, validate model training or check layout.")
print(json.dumps(result,indent=2))
if errors:raise SystemExit(1)
