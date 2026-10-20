#!/usr/bin/env python3
"""Check author-facing T1 Word revisions, without modifying them.
Usage: python code/analysis/verify_t1_integrated_word.py OLD_MAIN.docx NEW_MAIN.docx OLD_SUPP.docx NEW_SUPP.docx
Requirements: python-docx. Does not validate numerical truth of archived results.
"""
import sys,json,hashlib
from pathlib import Path
from docx import Document
if len(sys.argv)!=5:raise SystemExit(__doc__)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def cells(d):return [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
res=[]
for typ,(old,new) in zip(("main","supplement"),((Path(sys.argv[1]),Path(sys.argv[2])),(Path(sys.argv[3]),Path(sys.argv[4])))):
 a,b=Document(old),Document(new)
 body="\n".join(p.text for p in b.paragraphs)
 phrases=("T1 identity-control and historical-input provenance","Specific T1 tokenization and decision audit","Protocol-equivalence qualification for RQ4") if typ=="main" else ("Historical T1 input-serialization provenance","T1 tokenizer-count interpretation and limitations","T1 qualification to J2/J3")
 res.append({"document":typ,"old_sha256":sha(old),"new_sha256":sha(new),"old_tables":len(a.tables),"new_tables":len(b.tables),"all_table_cell_text_equal":cells(a)==cells(b),"all_three_contextual_sections_present":all(x in body for x in phrases)})
failed=any(not (r["all_table_cell_text_equal"] and r["all_three_contextual_sections_present"] and r["new_tables"]==(20 if r["document"]=="main" else 80)) for r in res)
print(json.dumps({"status":"FAIL" if failed else "PASS_EDITORIAL_INTEGRATION_ONLY","documents":res,"limit":"Not a check of empirical metrics, causal validity, or all visual formatting"},indent=2))
if failed:raise SystemExit(1)
