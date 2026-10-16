#!/usr/bin/env python3
"""Evidence-scoped read-only human reference integrity audit.

Requires: python >=3.10, openpyxl
Usage:
  python code/analysis/audit_human_reference_integrity_20261017.py \
    --gold /path/to/P2_FINAL_SENIOR_ADJUDICATED_GOLD_v1.0.xlsx \
    --audit-zip /path/to/HUMAN_REFERENCE_AUDIT_PACKAGE.zip \
    --output /tmp/human_reference_audit.json

This script does not execute an independent human validation or reconstruct
the original six-field routing decision. It never writes to GOLD.
"""
from __future__ import annotations
import argparse, collections, csv, hashlib, io, json, zipfile
from pathlib import Path
from openpyxl import load_workbook

KNOWN_GOLD_SHA = "ccc910b7bafbfa9607e93ef2eba8605f56f7ed87460ef74892d4877c7db4de65"

def digest(path):
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for chunk in iter(lambda:f.read(1048576),b""): h.update(chunk)
    return h.hexdigest()

def records(sheet):
    it=iter(sheet.values)
    header=next(it)
    return [dict(zip(header,row)) for row in it if row and row[0] is not None]

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--gold",required=True,type=Path)
    p.add_argument("--audit-zip",type=Path)
    p.add_argument("--output",type=Path,required=True)
    a=p.parse_args()
    audit={"scope":"gold workbook and optional derivative audit zip","external_human_validation":"NOT CONDUCTED","checks":{},"limitations":[]}
    def check(name,actual,expected):
        audit["checks"][name]={"observed":actual,"expected":expected,
           "status":"PASS" if actual==expected else "FAIL"}
    audit["gold_sha256"]=digest(a.gold)
    check("gold_sha256",audit["gold_sha256"],KNOWN_GOLD_SHA)
    wb=load_workbook(a.gold,read_only=True,data_only=True)
    needed={"FINAL_GOLD","TARGETED_REREVIEW","CONSISTENCY_AUDIT"}
    if not needed.issubset(wb.sheetnames):
        raise ValueError(f"Missing required sheets: {sorted(needed-set(wb.sheetnames))}")
    gold=records(wb["FINAL_GOLD"])
    rereview=records(wb["TARGETED_REREVIEW"])
    ids=[str(x["Case_ID"]) for x in gold]
    byid=dict(zip(ids,gold))
    check("gold_rows",len(gold),720)
    check("unique_case_ids",len(set(ids)),720)
    check("prompt_clusters",len({str(x["Prompt_ID"]) for x in gold}),120)
    check("final_failures",sum(str(x["Any_Failure"]).lower() in ("true","1","yes") for x in gold),474)
    check("rereview_rows",len(rereview),29)
    check("rereview_unique_ids",len({x["Case_ID"] for x in rereview}),29)
    check("rereview_ids_in_gold",sum(str(x["Case_ID"]) in byid for x in rereview),29)
    fields={
      "primary":("Final_Primary_After_Normalization","Final_Primary"),
      "secondary":("Final_Secondary_After_Normalization","Final_Secondary"),
      "boundary":("Final_Boundary","Final_Boundary"),
      "severity":("Final_Severity","Final_Severity"),
      "href_subtype":("Final_HREF_Subtype","Final_HREF_Subtype")
    }
    if all(str(x["Case_ID"]) in byid for x in rereview):
        for label,(source,target) in fields.items():
            check("review_final_"+label,sum(
                str(x[source])==str(byid[str(x["Case_ID"])][target])
                for x in rereview),29)
    check("review_revised",sum(x["Comparison_Outcome"]=="REVISE_TO_REREVIEW" for x in rereview),27)
    check("review_confirmed",sum(x["Comparison_Outcome"]!="REVISE_TO_REREVIEW" for x in rereview),2)
    audit["provenance_counts"]=dict(collections.Counter(str(x.get("Decision_Provenance")) for x in gold))
    if a.audit_zip:
        audit["audit_zip_sha256"]=digest(a.audit_zip)
        with zipfile.ZipFile(a.audit_zip) as z:
            changes=list(csv.DictReader(io.StringIO(z.read("TARGETED_29_FIELD_CHANGE_AUDIT.csv").decode("utf-8-sig"))))
        check("review_primary_changed",sum(str(x["Primary_Changed"]).strip().lower()=="yes" for x in changes),23)
        check("review_change_ids",set(str(x["Case_ID"]) for x in changes)==set(str(x["Case_ID"]) for x in rereview),True)
    audit["limitations"]=[
      "Original A1-A3 six-field and A4 decision sheets are NOT replayed by this script.",
      "The 297/423 historical routing criterion is NOT independently reconstructed.",
      "No independent recomputation of reliability interval estimates is asserted.",
      "Archive content matches do not establish construct or criterion validity.",
      "External human validation is deliberately future work."
    ]
    statuses=[x["status"] for x in audit["checks"].values()]
    audit["status"]="FAIL" if "FAIL" in statuses else "PASS_WITH_SCOPE_LIMITATIONS"
    a.output.parent.mkdir(parents=True,exist_ok=True)
    a.output.write_text(json.dumps(audit,indent=2,ensure_ascii=False,default=str)+"\n",encoding="utf-8")
    print(audit["status"],len(statuses),"checks")
    if "FAIL" in statuses:raise SystemExit(1)

if __name__=="__main__": main()
