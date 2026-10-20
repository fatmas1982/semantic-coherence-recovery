#!/usr/bin/env python3
"""Audit Word source table identities and numerical verification boundaries.
Usage:
  python code/analysis/audit_current_docx_table_inventory_20261020.py MAIN.docx SUPPLEMENT.docx --out audit_current_word
Requires python-docx. Does not certify full scientific numerical accuracy.
"""
import argparse,csv,hashlib,json,re
from pathlib import Path
from collections import Counter
from docx import Document
from docx.oxml.ns import qn
p=argparse.ArgumentParser()
p.add_argument("main",type=Path);p.add_argument("supplement",type=Path)
p.add_argument("--out",type=Path,default=Path("audit_current_word"))
args=p.parse_args();args.out.mkdir(parents=True,exist_ok=True)
SCOPED={
("MAIN","Table 10"):"SCOPED_SOURCE_DISPLAY_PASS",
("MAIN","Table 14"):"SCOPED_SOURCE_DISPLAY_PASS",
("MAIN","Table 17"):"SCOPED_SOURCE_DISPLAY_PASS",
("SUPPLEMENT","Table S14B"):"ARCHIVED_EXPORT_CONTRACT_PASS",
("SUPPLEMENT","Table S32A"):"ARCHIVED_EXPORT_CONTRACT_PASS",
("SUPPLEMENT","Table S37A"):"ARCHIVED_EXPORT_CONTRACT_PASS",
("SUPPLEMENT","Table S-E1E2"):"POINT_ESTIMATES_PASS_HISTORICAL_CI_OPEN",
("SUPPLEMENT","Table S44C"):"SCOPED_SOURCE_DISPLAY_PASS",
("SUPPLEMENT","Table S44E"):"SCOPED_SOURCE_DISPLAY_PASS"
}
entries=[]
for kind,file,expected in (("MAIN",args.main,20),("SUPPLEMENT",args.supplement,80)):
 d=Document(file)
 if len(d.tables)!=expected: raise SystemExit("Unexpected "+kind+" table count")
 queue=[];index=0
 for block in d.element.body.iterchildren():
  if block.tag==qn("w:p"):
   txt="".join(t.text or "" for t in block.iter(qn("w:t"))).strip()
   if txt:queue.append(txt);queue=queue[-12:]
  if block.tag!=qn("w:tbl"):continue
  index+=1;label="UNCAPTIONED_TABLE_"+str(index)
  for txt in reversed(queue):
   m=re.match(r"^(Table\s+(?:S(?:-E1E2|\d+[A-Z]?)|\d+[A-Z]?))\.\s*",txt)
   if m:label=m.group(1);break
  queue=[]
  t=d.tables[index-1]
  payload="|".join("||".join(c.text for c in row.cells) for row in t.rows)
  entries.append(dict(document=kind,table_index=index,published_caption=label,
   rows=len(t.rows),columns=max(len(z.cells) for z in t.rows),
   text_sha256=hashlib.sha256(payload.encode()).hexdigest(),
   scope_status=SCOPED.get((kind,label),"NO_NEW_FULL_TABLE_NUMERIC_CERTIFICATION"),
   source_word_sha256=hashlib.sha256(file.read_bytes()).hexdigest()))
assert len(entries)==100
with (args.out/"100_TABLE_DOCX_EVIDENCE_REGISTRY.csv").open("w",newline="",encoding="utf-8") as f:
 w=csv.DictWriter(f,fieldnames=list(entries[0]));w.writeheader();w.writerows(entries)
s={"inventory_count":len(entries),"by_status":dict(Counter(x["scope_status"] for x in entries)),
   "full_scientific_table_recomputations_verified_by_this_script":0,
   "important":"Displayed source metric matches are partial evidence, not confirmation of all table CIs, denominators and protocols"}
(args.out/"SUMMARY.json").write_text(json.dumps(s,indent=2)+"\n",encoding="utf-8")
print(json.dumps(s,indent=2))
